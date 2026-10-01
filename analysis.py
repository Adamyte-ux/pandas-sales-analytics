from pathlib import Path
import pandas as pd

# ============================================================
# Namaa Market - Sales Analytics Pipeline
# شركة نماء ماركت - مشروع تحليل بيانات تدريبي
# ============================================================

BASE_DIR = Path(__file__).resolve().parent
OUTPUT_FILE = BASE_DIR / "Namaa_Market_Analysis.xlsx"


def read_data():
    """قراءة ملفات الشركة الثلاثة."""
    customers = pd.read_csv(BASE_DIR / "customers.csv")
    products = pd.read_csv(BASE_DIR / "products.csv")
    orders = pd.read_csv(BASE_DIR / "orders.csv")
    return customers, products, orders


def check_reference_tables(customers, products):
    """التحقق من عدم تكرار مفاتيح العملاء والمنتجات."""
    if customers["Customer_ID"].duplicated().any():
        raise ValueError("يوجد Customer_ID مكرر في ملف العملاء")

    if products["Product_ID"].duplicated().any():
        raise ValueError("يوجد Product_ID مكرر في ملف المنتجات")


def clean_orders(orders):
    """تنظيف بيانات الطلبات مع الاحتفاظ بتقرير جودة البيانات."""
    quality = {
        "raw_rows": len(orders),
        "duplicate_rows": int(orders.duplicated().sum()),
        "missing_quantity": int(orders["Quantity"].isna().sum()),
        "missing_discount": int(orders["Discount_Pct"].isna().sum()),
    }

    cleaned = orders.drop_duplicates().copy()

    cleaned["Order_Date"] = pd.to_datetime(
        cleaned["Order_Date"], errors="coerce"
    )
    cleaned["Quantity"] = pd.to_numeric(
        cleaned["Quantity"], errors="coerce"
    ).fillna(1)
    cleaned["Discount_Pct"] = pd.to_numeric(
        cleaned["Discount_Pct"], errors="coerce"
    ).fillna(0)

    # حماية من القيم غير المنطقية
    cleaned.loc[cleaned["Quantity"] <= 0, "Quantity"] = 1
    cleaned.loc[cleaned["Discount_Pct"] < 0, "Discount_Pct"] = 0
    cleaned.loc[cleaned["Discount_Pct"] > 1, "Discount_Pct"] = 1

    quality["cleaned_rows"] = len(cleaned)
    quality["invalid_dates"] = int(cleaned["Order_Date"].isna().sum())
    quality["missing_values_after_cleaning"] = int(cleaned.isna().sum().sum())

    return cleaned, pd.DataFrame(
        [{"Metric": key, "Value": value} for key, value in quality.items()]
    )


def enrich_orders(orders, customers, products):
    """دمج الطلبات مع بيانات المنتجات والعملاء وحساب المؤشرات المالية."""
    enriched = orders.merge(
        products,
        on="Product_ID",
        how="left",
        validate="many_to_one",
        indicator="Product_Match",
    )

    enriched = enriched.merge(
        customers,
        on="Customer_ID",
        how="left",
        validate="many_to_one",
        indicator="Customer_Match",
    )

    missing_products = int((enriched["Product_Match"] != "both").sum())
    missing_customers = int((enriched["Customer_Match"] != "both").sum())

    if missing_products:
        print(f"تحذير: {missing_products} طلبات بلا منتج مطابق")
    if missing_customers:
        print(f"تحذير: {missing_customers} طلبات بلا عميل مطابق")

    enriched["Gross_Revenue"] = (
        enriched["Unit_Price"] * enriched["Quantity"]
    )
    enriched["Discount_Amount"] = (
        enriched["Gross_Revenue"] * enriched["Discount_Pct"]
    )
    enriched["Net_Revenue"] = (
        enriched["Gross_Revenue"] - enriched["Discount_Amount"]
    )
    enriched["Total_Cost"] = (
        enriched["Unit_Cost"] * enriched["Quantity"]
    )
    enriched["Gross_Profit"] = (
        enriched["Net_Revenue"] - enriched["Total_Cost"]
    )
    enriched["Profit_Margin"] = enriched["Gross_Profit"] / enriched["Net_Revenue"].where(
        enriched["Net_Revenue"] != 0
    )
    enriched["Month"] = enriched["Order_Date"].dt.to_period("M").astype(str)
    enriched["Is_Realized_Revenue"] = enriched["Order_Status"].eq("Delivered")

    enriched = enriched.drop(columns=["Product_Match", "Customer_Match"])
    return enriched


def build_analyses(enriched):
    """إنشاء جداول التحليل التي ستذهب إلى أوراق Excel."""
    delivered = enriched[enriched["Order_Status"] == "Delivered"].copy()

    summary = pd.DataFrame([
        ["All order rows", len(enriched)],
        ["Delivered orders", int((enriched["Order_Status"] == "Delivered").sum())],
        ["Pending orders", int((enriched["Order_Status"] == "Pending").sum())],
        ["Cancelled orders", int((enriched["Order_Status"] == "Cancelled").sum())],
        ["Returned orders", int((enriched["Order_Status"] == "Returned").sum())],
        ["Realized net revenue", delivered["Net_Revenue"].sum()],
        ["Realized gross profit", delivered["Gross_Profit"].sum()],
        ["Average delivered order value", delivered["Net_Revenue"].mean()],
    ], columns=["Metric", "Value"])

    by_city = (
        delivered.groupby("City", as_index=False)
        .agg(
            Orders=("Order_ID", "nunique"),
            Quantity=("Quantity", "sum"),
            Revenue=("Net_Revenue", "sum"),
            Profit=("Gross_Profit", "sum"),
        )
        .sort_values("Revenue", ascending=False)
    )
    by_city["Profit_Margin"] = by_city["Profit"] / by_city["Revenue"]

    by_month = (
        delivered.groupby("Month", as_index=False)
        .agg(
            Orders=("Order_ID", "nunique"),
            Revenue=("Net_Revenue", "sum"),
            Profit=("Gross_Profit", "sum"),
        )
        .sort_values("Month")
    )
    by_month["Profit_Margin"] = by_month["Profit"] / by_month["Revenue"]

    by_product = (
        delivered.groupby(["Product_ID", "Product_Name", "Category"], as_index=False)
        .agg(
            Orders=("Order_ID", "nunique"),
            Quantity=("Quantity", "sum"),
            Revenue=("Net_Revenue", "sum"),
            Profit=("Gross_Profit", "sum"),
        )
        .sort_values("Revenue", ascending=False)
    )
    by_product["Profit_Margin"] = by_product["Profit"] / by_product["Revenue"]

    by_category = (
        delivered.groupby("Category", as_index=False)
        .agg(
            Orders=("Order_ID", "nunique"),
            Quantity=("Quantity", "sum"),
            Revenue=("Net_Revenue", "sum"),
            Profit=("Gross_Profit", "sum"),
        )
        .sort_values("Revenue", ascending=False)
    )
    by_category["Profit_Margin"] = by_category["Profit"] / by_category["Revenue"]

    by_status = (
        enriched.groupby("Order_Status", as_index=False)
        .agg(
            Orders=("Order_ID", "nunique"),
            Rows=("Order_ID", "count"),
            Gross_Revenue=("Gross_Revenue", "sum"),
            Net_Revenue=("Net_Revenue", "sum"),
        )
        .sort_values("Orders", ascending=False)
    )

    by_payment = (
        delivered.groupby("Payment_Method", as_index=False)
        .agg(
            Orders=("Order_ID", "nunique"),
            Revenue=("Net_Revenue", "sum"),
            Profit=("Gross_Profit", "sum"),
        )
        .sort_values("Revenue", ascending=False)
    )

    top_customers = (
        delivered.groupby(["Customer_ID", "Customer_Name", "City", "Segment"], as_index=False)
        .agg(
            Orders=("Order_ID", "nunique"),
            Quantity=("Quantity", "sum"),
            Revenue=("Net_Revenue", "sum"),
            Profit=("Gross_Profit", "sum"),
        )
        .sort_values("Revenue", ascending=False)
    )

    return {
        "Summary": summary,
        "By_City": by_city,
        "By_Month": by_month,
        "By_Product": by_product,
        "By_Category": by_category,
        "By_Status": by_status,
        "By_Payment": by_payment,
        "Top_Customers": top_customers,
    }


def format_workbook(writer):
    """تنسيق أوراق Excel لتكون أسهل في القراءة."""
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    for worksheet in writer.book.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center")

        for column_cells in worksheet.columns:
            max_length = max(len(str(cell.value or "")) for cell in column_cells)
            col_letter = get_column_letter(column_cells[0].column)
            worksheet.column_dimensions[col_letter].width = min(max(max_length + 2, 12), 28)

        for row in worksheet.iter_rows():
            for cell in row:
                if cell.column >= 1 and isinstance(cell.value, float):
                    cell.number_format = '#,##0.00'


def export_report(customers, products, raw_orders, cleaned_orders, enriched, quality, analyses):
    """تصدير كل البيانات والجداول إلى ملف Excel واحد."""
    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        customers.to_excel(writer, sheet_name="Customers", index=False)
        products.to_excel(writer, sheet_name="Products", index=False)
        raw_orders.to_excel(writer, sheet_name="Raw_Orders", index=False)
        cleaned_orders.to_excel(writer, sheet_name="Clean_Orders", index=False)
        enriched.to_excel(writer, sheet_name="Enriched_Orders", index=False)
        quality.to_excel(writer, sheet_name="Data_Quality", index=False)

        for sheet_name, dataframe in analyses.items():
            dataframe.to_excel(writer, sheet_name=sheet_name, index=False)

        format_workbook(writer)


def main():
    customers, products, raw_orders = read_data()
    check_reference_tables(customers, products)

    cleaned_orders, quality = clean_orders(raw_orders)
    enriched_orders = enrich_orders(cleaned_orders, customers, products)
    analyses = build_analyses(enriched_orders)

    export_report(
        customers,
        products,
        raw_orders,
        cleaned_orders,
        enriched_orders,
        quality,
        analyses,
    )

    print("تم إنشاء التقرير بنجاح:")
    print(OUTPUT_FILE)
    print("\nملخص سريع:")
    print(analyses["Summary"].to_string(index=False))


if __name__ == "__main__":
    main()
